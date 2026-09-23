from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "DAS7002_Final_Report_Polished.docx"
CHARTS = ROOT / "outputs" / "charts"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_run_font(run, size=10)


def suppress_paragraph_borders(paragraph):
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    paragraph_properties.append(borders)


def add_paragraph(doc, text, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_after = Pt(space_after)
    if first_line:
        fmt.first_line_indent = Cm(1.27)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        set_run_font(run, size=14 if level == 1 else 13, bold=True)
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=10, italic=True)
    return paragraph


def add_figure(doc, filename, caption, width=6.1):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    picture = paragraph.add_run().add_picture(str(CHARTS / filename), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(caption)
    set_run_font(run, size=10, italic=True)


def add_table(doc, headers, rows, widths, caption):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_borders(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(header))
        set_run_font(run, size=9, bold=True)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "EAF2F8")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_run_font(run, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.8)
    paragraph.paragraph_format.first_line_indent = Cm(-0.4)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(text))


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(header.add_run("DAS 7002 Big Data Technologies"), size=10)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(90)
    title.paragraph_format.space_after = Pt(24)
    run = title.add_run("Distributed Analysis of Chicago Crime Data Using Apache Spark")
    set_run_font(run, size=20, bold=True)
    run.font.color.rgb = RGBColor(0, 0, 0)
    suppress_paragraph_borders(title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("Data Engineering Exploratory Analysis Clustering and Arrest Prediction"), size=14, bold=True)
    subtitle.paragraph_format.space_after = Pt(48)

    details = [
        ("Module", "DAS 7002 Big Data Technologies"),
        ("Assessment", "PRAC 1 Data Analysis Using Big Data and Distributed Computing"),
        ("Student name", "[Enter student name]"),
        ("Student ID", "[Enter student ID]"),
        ("Submission date", "[Enter submission date]"),
        ("Word count", "[Updated after final word-count check]"),
    ]
    table = doc.add_table(rows=len(details), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (label, value) in zip(table.rows, details):
        row.cells[0].width = Inches(1.7)
        row.cells[1].width = Inches(4.7)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell, color="FFFFFF", size="0")
        set_run_font(row.cells[0].paragraphs[0].add_run(label), bold=True)
        set_run_font(row.cells[1].paragraphs[0].add_run(value))

    declaration = doc.add_paragraph()
    declaration.paragraph_format.space_before = Pt(55)
    declaration.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run_font(declaration.add_run("Declaration  "), bold=True)
    set_run_font(declaration.add_run(
        "I confirm that this submission is my own work and that all external ideas, data and sources are acknowledged through Harvard-style citations and the reference list."
    ), size=11)

    doc.add_page_break()
    add_heading(doc, "Abstract", 1)
    add_paragraph(doc,
        "This report presents a distributed analysis of the Chicago Crime dataset using Apache Spark. The study covers four connected stages: data engineering, exploratory analysis, spatial-temporal clustering and supervised prediction of arrest outcomes. The raw crime file contained 7,784,664 records. An explicit schema and quality policy converted source strings to typed fields, rejected invalid identifiers or timestamps, retained imperfect geographic records with validity flags, and stored the cleaned data as partitioned Parquet. Socioeconomic and hourly weather data were processed as secondary sources. PySpark SQL identified temporal and community patterns, calculated seven-day rolling averages, and compared weather conditions with crime frequency. Weather correlations were interpreted as associations rather than causal effects. K-means models with K values from two to six were repeated across three seeds. K equals two produced the strongest mean silhouette of 0.496 with a standard deviation of 0.00035, although its broad zones crossed every police district. Weighted Logistic Regression achieved the highest F1 score of 0.685, while Gradient Boosted Trees achieved the highest recall of 0.582 and ROC AUC of 0.848. A temporal Random Forest holdout produced a lower F1 score of 0.493, showing that random validation gave an optimistic estimate of future performance. Spark therefore supports a reproducible end-to-end workflow at this scale, but reported crime data, airport weather observations and police outcomes limit interpretation and operational use."
    )

    add_heading(doc, "Contents", 1)
    contents = [
        "1 Introduction", "2 Data and distributed architecture", "3 Task 1 Data engineering and pipeline construction",
        "4 Task 2 Exploratory analysis", "5 Task 3 Spatial and temporal clustering",
        "6 Task 4 Distributed arrest prediction", "7 Performance reproducibility and software quality",
        "8 Ethical issues and study limitations", "9 Conclusion and recommendations", "References", "Appendix A Reproducibility commands",
    ]
    for item in contents:
        add_paragraph(doc, item, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)

    add_heading(doc, "1 Introduction", 1)
    add_paragraph(doc,
        "Large public safety datasets create a practical need for systems that can store, transform and analyse millions of records without relying on manual spreadsheet work. The Chicago Crime dataset is suitable for this purpose because it contains reported incidents from 2001 onwards and includes timestamps, police districts, community areas, coordinates, crime categories and arrest outcomes. The City of Chicago explains that these are reported incidents extracted from the police CLEAR system, that recent records are withheld for seven days, and that classifications can change after further investigation (City of Chicago, 2026). These conditions matter because the dataset is large and detailed, but it is not a complete or final account of all crime in the city. The analysis must therefore combine scalable computation with careful interpretation."
    )
    add_paragraph(doc,
        "The purpose of this project is to build and evaluate one coherent Spark workflow across the four assessment tasks. Task 1 creates an auditable extract, transform and load pipeline for crime, census and weather data. Task 2 uses PySpark SQL and DataFrame operations to examine time, location, socioeconomic context and weather. Task 3 applies K-means to latitude, longitude and time of day, then tests whether activity zones match police districts. Task 4 compares distributed Random Forest, Logistic Regression and Gradient Boosted Tree classifiers for arrest prediction and examines class imbalance. Spark provides an effective technical foundation, but reliable conclusions depend on explicit quality rules, appropriate validation and an understanding of how the source data were produced."
    )
    add_paragraph(doc,
        "The report follows the same sequence as the implemented pipeline. This structure keeps the logical connection between decisions and evidence visible. Clean fields and quality flags from Task 1 support the rolling windows and joins in Task 2. Valid coordinates and hours support clustering in Task 3. Typed temporal, geographic and categorical fields become model features in Task 4. This dependency is important: a prediction score cannot repair an incorrect timestamp, and a cluster cannot represent geography when its coordinates are missing. Each section therefore explains the implementation, gives measured results, and then evaluates limitations before moving to the next task."
    )
    add_paragraph(doc,
        "The work was executed locally with PySpark 3.5.9 using four worker threads and a two-gigabyte driver heap. This is a Spark environment, but it is not a multi-node production cluster. Distributed transformations, shuffles, ML pipelines and Parquet partitioning are genuine Spark operations. However, the measured run times and memory behaviour do not demonstrate horizontal cluster scalability. The design remains transferable because input and output paths are configurable, transformations are separated into reusable modules, dependency versions are fixed, and the same commands can run against a larger Spark deployment with appropriate storage and executor settings."
    )

    add_heading(doc, "2 Data and distributed architecture", 1)
    add_heading(doc, "2.1 Data sources", 2)
    add_paragraph(doc,
        "The primary source is the Chicago Crimes 2001 to Present dataset. It records reported incidents rather than every offence that occurred, and most addresses are reduced to block level to protect victims (City of Chicago, 2026). The local CSV used in this project contained 7,784,664 records. Its scale makes repeated full file scans expensive and justifies conversion from CSV to a typed columnar format. It also contains several levels of geography. Police district is an administrative field, community area supports neighbourhood analysis, and latitude and longitude support spatial modelling. These levels are related but not interchangeable, so the pipeline keeps all three rather than replacing one with another."
    )
    add_paragraph(doc,
        "The socioeconomic file contains selected indicators for Chicago community areas during 2008 to 2012. The official description identifies six indicators and a composite hardship index derived for public health analysis (Chicago Department of Public Health, 2014). The project uses community number, community name and hardship index as a small reference table. The weather file contains 62,269 observations from one Chicago Midway Airport station over the same period. NOAA Local Climatological Data includes hourly temperature, precipitation, humidity, wind and weather type, while NOAA also warns that stations can have different variables and gaps (Kantor et al., 2023). Midway is therefore treated as a city-wide weather proxy, not as a measurement at every crime location."
    )
    add_table(doc,
        ["Source", "Raw scale", "Analytical role", "Main limitation"],
        [
            ["Chicago crime", "7,784,664 records", "Core incident, spatial and arrest analysis", "Reported police incidents are not a census of crime"],
            ["Socioeconomic", "77 community areas", "Names and hardship context", "Indicators represent 2008 to 2012 only"],
            ["NOAA weather", "62,269 observations", "Hourly and daily weather comparison", "One airport station represents the wider city"],
        ],
        [1.25, 1.15, 2.15, 2.3],
        "Table 1 Data sources and analytical roles",
    )

    add_heading(doc, "2.2 Architecture and design choices", 2)
    add_paragraph(doc,
        "The architecture separates command-line entry points from transformation functions. The Task 1 runner reads paths and constructs an immutable configuration object. The orchestration layer creates the Spark session, validates required columns, calls pure DataFrame transformations and writes outputs. Separate modules handle census cleaning, weather aggregation, exploratory analysis, clustering, modelling, charts and evidence export. This division applies the single-responsibility principle at module level and reduces duplication. For example, the weather-condition expression is defined once and reused in overall and seasonal comparisons. The same Random Forest pipeline builder is used for weighted, unweighted and temporal models, preventing small feature differences from weakening the comparison."
    )
    add_paragraph(doc,
        "Spark DataFrames were preferred to low-level RDD operations because the tasks are structured and depend on typed columns, SQL aggregation, joins, window functions and Spark ML. Spark's unified engine supports different analytical workloads through a shared programming model (Zaharia et al., 2016). Spark SQL can optimise logical plans and work with several data sources through the DataFrame interface (Apache Software Foundation, 2026c). Built-in expressions were used instead of Python user-defined functions wherever possible. This keeps computation inside Spark's optimised execution engine and makes column lineage easier to inspect. Temporary SQL views were retained where the assessment explicitly required PySpark SQL, particularly for the temporal pattern query."
    )
    add_paragraph(doc,
        "The pipeline stores cleaned data in Parquet. Parquet is column-oriented and supports efficient storage and retrieval through compression and encoding for large analytical datasets (Apache Parquet, 2026). This format is more suitable than repeatedly parsing CSV because later tasks read selected typed columns rather than every original string. Crime output is partitioned by year and district. The layout supports partition pruning for common filters, such as analysis of one year or district. It can also produce many small files when a dataset has numerous combinations, so the design balances query selectivity against file management. The current local output remained practical, while a production data lake would require periodic compaction and monitoring of partition sizes."
    )
    add_paragraph(doc,
        "The Spark session uses adaptive query execution and eight shuffle partitions in the local environment. Large reusable frames are persisted to disk when an algorithm performs several actions. Disk persistence was chosen because the two-gigabyte driver and local memory limit cannot safely hold all transformed records. During the full modelling run, Spark reported that some blocks were spilled to disk. The run still completed, confirming that the persistence strategy protected correctness, although it increased execution time. In a cluster deployment, executor memory, partition counts and storage levels should be tuned from measured workload statistics rather than copied from these local settings."
    )

    add_heading(doc, "3 Task 1 Data engineering and pipeline construction", 1)
    add_heading(doc, "3.1 Schema enforcement", 2)
    add_paragraph(doc,
        "The raw crime CSV is first read as strings. This preserves malformed values until the pipeline can distinguish a parsing error from a genuinely missing source value. After header normalisation, the program checks for nine required fields: ID, Date, District, Community Area, Latitude, Longitude, Arrest, Domestic and Primary Type. A missing required column causes an immediate error. The pipeline then applies explicit casts, supported timestamp patterns, trimmed text and lowercase Boolean conversion. Derived columns include year, month, day of week and hour, which are reused by later tasks."
    )
    add_paragraph(doc,
        "The source to target data dictionary records the original field, target name, Spark type, cleaning rule and invalid value policy. Crime ID becomes a long integer and must be positive. Date becomes a timestamp and must match an accepted format. District and community area become integers, while coordinates become doubles. Arrest and Domestic become Boolean values, and Primary Type becomes trimmed text. Keeping this policy as an output table makes the transformation auditable. It also separates technical type enforcement from analytical judgement. A field can be correctly converted to an integer but still fall outside a valid Chicago range, so casting and validation are treated as separate steps."
    )
    add_table(doc,
        ["Field", "Target type", "Quality rule", "Action"],
        [
            ["ID", "Long", "Positive and not null", "Reject invalid record"],
            ["Date", "Timestamp", "Supported format and not null", "Reject invalid record"],
            ["District", "Integer", "Range 1 to 31", "Retain with validity flag"],
            ["Community Area", "Integer", "Range 1 to 77", "Retain with validity flag"],
            ["Latitude and Longitude", "Double", "Within configured Chicago bounds", "Retain with validity flag"],
            ["Arrest and Domestic", "Boolean", "Normalised true or false", "Retain null and filter for modelling"],
        ],
        [1.55, 1.15, 2.35, 1.8],
        "Table 2 Core schema and invalid value policy",
    )

    add_heading(doc, "3.2 Data quality policy", 2)
    add_paragraph(doc,
        "Records with an invalid crime ID or timestamp are quarantined because they cannot be reliably identified, deduplicated or placed in time. The quarantine dataset includes a rejection reason so that failures can be reviewed without losing the source row. Geographic errors receive a different treatment. A missing coordinate prevents mapping and clustering, but the same incident may still be valid for city-wide counts, crime-type analysis or arrest modelling if other fields are present. Deleting every row with a geographic defect would reduce the dataset and could introduce additional selection bias. The pipeline therefore retains these records and adds flags for valid coordinates, districts and community areas. Downstream tasks apply only the flags they need."
    )
    add_paragraph(doc,
        "The quality summary found 7,784,664 raw records and the same number of distinct cleaned crime IDs. There were no invalid IDs or timestamps in this extract, so the quarantine output was empty apart from its schema and metadata. Geographic checks found 86,966 records with invalid coordinates, 47 with invalid districts and 613,552 with invalid community areas. These categories can overlap and should not be added as if they were mutually exclusive. The outcome shows why quality reporting remains necessary even when no rows are rejected. A pipeline that only reports final row count would hide the fact that several hundred thousand incidents cannot support every spatial operation."
    )
    add_table(doc,
        ["Measure", "Count", "Share of raw records"],
        [
            ["Raw and cleaned records", "7,784,664", "100.00%"],
            ["Invalid crime IDs", "0", "0.00%"],
            ["Invalid timestamps", "0", "0.00%"],
            ["Invalid coordinates", "86,966", "1.12%"],
            ["Invalid districts", "47", "Less than 0.01%"],
            ["Invalid community areas", "613,552", "7.88%"],
        ],
        [3.1, 1.6, 2.1],
        "Table 3 Crime data quality results",
    )
    add_paragraph(doc,
        "Deduplication is based on crime ID after type conversion and mandatory validation. This ordering is important because duplicate null identifiers would otherwise collapse unrelated records. In the current file, the cleaned and distinct counts were equal, but the rule remains necessary for future extracts. The quality calculations are produced both before and after cleaning, then stored as Parquet and exported to CSV for inspection. Structured logging identifies each stage and output path. These choices make data loss visible and reproducible rather than leaving it as an undocumented side effect of filtering."
    )

    add_heading(doc, "3.3 Secondary data and partitioned output", 2)
    add_paragraph(doc,
        "The socioeconomic cleaner standardises the community area key, community name and hardship index, removes the citywide total row, and produces 77 valid community areas. This number matches the recognised community area structure and gives a compact lookup table for broadcast joins. Weather processing parses the observation timestamp, converts numerical measures, recognises trace precipitation, and derives event flags for rain, snow and thunderstorms. Multiple observations within an hour are aggregated to one station hour. The process reduced 62,269 raw weather rows to 43,843 hourly rows covering 1 January 2008 to 31 December 2012. No invalid weather timestamps were found."
    )
    add_paragraph(doc,
        "The final crime output is written with overwrite or fail-if-exists behaviour controlled through the command line. Partition columns are year and district, while the remaining columns stay within Parquet files. Spark can discover partition values from directory paths when reading partitioned data (Apache Software Foundation, 2026c). Later code can therefore read a normal DataFrame without reconstructing year and district manually. The approach also supports selective processing; for example, a query restricted to 2012 and one district can avoid scanning unrelated partitions. Partitioning by community area could help neighbourhood queries, but year and district better match the assessment requirement and avoid another high-cardinality directory level."
    )

    add_heading(doc, "4 Task 2 Exploratory analysis", 1)
    add_heading(doc, "4.1 Temporal analysis and rolling averages", 2)
    add_paragraph(doc,
        "Temporal patterns were created with an explicit Spark SQL query that groups valid timestamps by year, month, day of week and hour. The hourly chart shows clear variation through the day rather than a constant incident rate. Weekday counts are lowest in the early morning, rise through the morning and remain higher from midday into late evening. Midnight also contains a strong peak. That value requires caution because some source systems use midnight when the exact time is unknown; the chart therefore describes recorded timestamps, not necessarily the precise time when every incident occurred. A separate weekday and weekend table allows the shape of the day to be compared without mixing different exposure patterns."
    )
    add_figure(doc, "task2_hourly_crime.png", "Figure 1 Crime frequency by recorded hour")
    add_paragraph(doc,
        "Community rolling averages were calculated over a complete date calendar rather than only dates that contained crime. For each of the 77 community areas, the pipeline crosses the area key with every weather date, left joins the daily crime count and fills missing counts with zero. A Spark window then partitions by community area, orders by date and averages the current day with the previous six rows. Window functions support moving averages and other calculations over related rows (Apache Software Foundation, 2026e). Building the complete calendar prevents a seven-row window from silently covering more than seven days when zero-crime dates are absent."
    )
    add_paragraph(doc,
        "The rolling result is enriched with community name and hardship index through a broadcast join. The census table contains only 77 rows, so sending it to each executor is cheaper than shuffling millions of crime records by community key. Spark documentation explains that a broadcast hint prioritises a broadcast join for the selected build side when the join type allows it (Apache Software Foundation, 2026d). The project applies this optimisation explicitly, satisfying the task and demonstrating an appropriate join between a large fact table and a small lookup table. The same technique is used for community summaries, while the large crime and weather aggregations remain distributed."
    )

    add_heading(doc, "4.2 Spatial and socioeconomic patterns", 2)
    add_paragraph(doc,
        "The community analysis is restricted to the 2008 to 2012 overlap among crime, weather and socioeconomic data. Austin recorded 120,692 incidents, representing 6.48 per cent of valid geocoded incidents in this period. South Shore recorded 62,901, Humboldt Park 59,659 and Near North Side 59,481. These results show why a socioeconomic explanation cannot be reduced to hardship alone. West Englewood and Englewood had high hardship indices of 89 and 94 and high crime counts, but Near North Side also had a high count with a hardship index of 1. Population, land use, visitor numbers, transport hubs and reporting behaviour are not controlled here. Raw incident counts should therefore not be interpreted as individual risk or as a direct effect of hardship."
    )
    add_paragraph(doc,
        "Each community summary includes a crime share and mean latitude and longitude. These centroids support a compact spatial comparison, but they do not reproduce official boundaries or show variation within a community. A choropleth based on boundary polygons would give a clearer representation of area rates, especially if crime counts were divided by population. The present evidence remains useful because it identifies concentration and provides a reproducible table for further mapping. It also keeps the distinction between administrative areas and point coordinates visible, which becomes central in Task 3."
    )

    add_heading(doc, "4.3 Weather associations", 2)
    add_paragraph(doc,
        "Crime and weather were first aligned at hourly grain using date and hour, then aggregated to daily records. This avoids joining every incident directly to every weather observation and reduces the data before statistical comparison. Event categories follow a fixed precedence: thunderstorm, snow, rain, other precipitation, other event and no event. The ordering prevents one day from appearing in several categories when its observations contain more than one code. Daily counts were compared through means and medians, and Pearson correlations were calculated for temperature, precipitation, humidity and wind."
    )
    add_figure(doc, "task2_weather_impact.png", "Figure 2 Average daily crime count by observed weather condition")
    add_paragraph(doc,
        "The strongest simple correlation was between average temperature and daily crime count at 0.510. Precipitation had a correlation of minus 0.028, humidity minus 0.176 and wind speed minus 0.151. On no-event days, the mean daily count was 1,049.3. Rain days were 2.0 per cent lower, snow days 12.8 per cent lower, and thunderstorm days 1.7 per cent higher. Other precipitation days were 2.3 per cent higher. These effect sizes are more interpretable than correlation coefficients alone because they state the size and direction of the observed difference. They remain descriptive because weather conditions are not randomly assigned, and day type, season, long-term change and human activity can influence both crime and weather exposure."
    )
    add_table(doc,
        ["Condition", "Days", "Mean daily crime", "Difference from no event"],
        [
            ["No event", "746", "1,049.3", "Reference"],
            ["Rain", "374", "1,028.6", "minus 2.0%"],
            ["Snow", "278", "915.0", "minus 12.8%"],
            ["Thunderstorm", "198", "1,067.4", "plus 1.7%"],
            ["Other precipitation", "112", "1,073.6", "plus 2.3%"],
        ],
        [2.0, 0.8, 1.65, 2.35],
        "Table 4 Daily crime counts by weather condition",
    )
    add_paragraph(doc,
        "Season controlled summaries demonstrate the confounding problem. The average on no event days was 1,123.2 in summer but only 899.4 in winter. Within summer, rain days averaged 1,125.2, which was almost the same as no event days. Within winter, snow days averaged 896.2, also close to the no event winter average. The large overall gap for snow therefore partly reflects the lower winter baseline rather than an isolated snow effect. Spring rain and thunderstorms were below the spring no event average, but the groups also had different numbers of days and substantial standard deviations. Formal causal analysis would require a regression or quasi experimental design with season, year, weekday, holiday and exposure controls."
    )
    add_paragraph(doc,
        "The weather results are also limited by measurement location. Midway Airport gives consistent hourly coverage, but Chicago covers a wide area and local precipitation can vary. Assigning the same station value to every community creates exposure measurement error. The method is defensible for a city-wide exploratory task because it is transparent and reproducible, yet it cannot support claims about neighbourhood-level weather effects. A stronger extension would combine several stations, choose the nearest station to each incident, or use gridded weather data. The report therefore describes associations and avoids causal claims."
    )

    add_heading(doc, "5 Task 3 Spatial and temporal clustering", 1)
    add_heading(doc, "5.1 Feature preparation and model selection", 2)
    add_paragraph(doc,
        "K-means clustering used latitude, longitude and recorded hour. Rows without valid coordinates or time were excluded because distance calculations require complete numerical vectors. VectorAssembler created a raw vector and StandardScaler transformed each dimension to zero-centred, unit-scale values. Without scaling, hour or a coordinate range could dominate squared Euclidean distance for numerical rather than substantive reasons. Model selection used a reproducible 0.2 per cent training sample, coalesced to eight partitions. Each K from two to six was fitted with seeds 42, 123 and 2026 before the selected model assigned every valid record. Reusing one sample isolates sensitivity to K-means initialisation rather than adding sampling variation."
    )
    add_paragraph(doc,
        "Spark's K-means implementation supports the scalable k-means parallel initialisation developed for large datasets (Apache Software Foundation, 2026b; Bahmani et al., 2012). Each candidate was evaluated through squared Euclidean silhouette and model training cost, which represents within-cluster squared error. The silhouette compares cohesion within a cluster with separation from other clusters; values nearer one indicate better separation (Rousseeuw, 1987). Training cost tends to fall as K increases because more centres can fit the data more closely. It must therefore be interpreted through the elbow shape rather than used to select the largest K automatically."
    )
    add_figure(doc, "task3_silhouette.png", "Figure 3 Mean silhouette score with standard deviation across three seeds")
    add_figure(doc, "task3_elbow.png", "Figure 4 Within cluster training cost for candidate K values")
    add_table(doc,
        ["K", "Mean silhouette (SD)", "Mean training cost", "Average dominant district share"],
        [
            ["2", "0.496 (0.00035)", "29,474.5", "0.125"],
            ["3", "0.490 (0.00042)", "22,346.1", "0.126"],
            ["4", "0.442 (0.00050)", "18,447.9", "0.129"],
            ["5", "0.414 (0.02489)", "16,638.2", "0.180"],
            ["6", "0.433 (0.00708)", "13,956.1", "0.203"],
        ],
        [0.55, 1.65, 1.55, 2.95],
        "Table 5 K-means validation evidence",
    )
    add_paragraph(doc,
        "K equals two achieved the highest mean silhouette of 0.496, narrowly above 0.490 for K equals three. Their standard deviations were only 0.00035 and 0.00042, demonstrating stable initialisation. K equals five was less stable, with a standard deviation of 0.02489. Training cost declined as K increased, creating a genuine trade-off between statistical separation and operational detail. The average dominant district share increased from 0.125 at K equals two to 0.203 at K equals six. Even at six clusters, one police district did not dominate a typical cluster, so the learned spatial-temporal structure does not follow official boundaries closely."
    )

    add_heading(doc, "5.2 Cluster interpretation", 2)
    add_figure(doc, "task3_cluster_map.png", "Figure 5 Sampled spatial distribution of the selected K-means clusters")
    add_paragraph(doc,
        "The selected model assigned 3,568,524 incidents to cluster 0 and 4,129,174 to cluster 1. Their mean hours were 13.20 and 13.11, so time of day did not strongly distinguish the final two groups. Their centroids mainly separate the city geographically: cluster 0 is centred near latitude 41.767 and longitude minus 87.629, while cluster 1 is centred near latitude 41.907 and longitude minus 87.708. Both clusters contain incidents from 22 districts and cross district boundaries. The activity zones therefore do not align with official police districts. However, they are broad regional divisions rather than precise hotspots."
    )
    add_paragraph(doc,
        "High activity was defined through an activity concentration index rather than cluster size alone. The index divides incident count by spatial dispersion, where dispersion is derived from latitude and longitude variance. Cluster 1 had an index of 59.13 million compared with 55.93 million for cluster 0 and was above the mean concentration. This definition recognises that a large cluster covering a very wide space is not necessarily a concentrated hotspot. It remains a relative internal measure, not a crime rate. The index does not adjust for population, land area, visitors or duration, and its numerical scale depends on coordinate units. It should therefore support comparison between these fitted clusters, not comparison with another city or model without standardisation."
    )
    add_paragraph(doc,
        "A decision maker might prefer K equals four or six because smaller zones can be described and resourced more easily. That choice would sacrifice silhouette quality but provide greater district concentration and local detail. K equals two is statistically defensible, but not automatically the best operational map. The three seeded runs increase confidence that its separation is not an initialisation accident. A stronger study would test stability across time and samples, compare K-means with density-based spatial methods, and validate clusters against external outcomes. It could also encode hour cyclically as sine and cosine because 23 and 0 are adjacent in time."
    )

    add_heading(doc, "6 Task 4 Distributed arrest prediction", 1)
    add_heading(doc, "6.1 Target features and imbalance", 2)
    add_paragraph(doc,
        "Task 4 predicts whether an incident resulted in an arrest. This target was selected instead of Primary Type because arrest is already a binary field and supports focused evaluation through precision, recall, F1, ROC AUC and PR AUC. Input features are latitude, longitude, hour, month, district, community area and Primary Type. StringIndexer converts Primary Type to a numeric index, VectorAssembler creates the tree-model feature vector, and RandomForestClassifier fits twenty trees with a fixed seed. Logistic Regression uses a one-hot-encoded version of Primary Type. The modelling data excludes rows with invalid coordinates, districts or community areas and rows with missing targets or crime types, leaving 7,093,315 records."
    )
    add_paragraph(doc,
        "The target is imbalanced. Non-arrest records account for 5,258,641 cases or 74.14 per cent, while arrest records account for 1,834,674 cases or 25.86 per cent. Class imbalance can make overall accuracy misleading and can reduce a learner's attention to the minority class (He and Garcia, 2009). A classifier that always predicts no arrest would therefore appear accurate on most records but would identify no arrests. The project records this majority baseline explicitly. The baseline has zero precision, recall and F1 for the positive class, balanced accuracy of 0.5, ROC AUC of 0.5 and PR AUC equal to the positive prevalence of 0.258. This comparison prevents the model from receiving credit merely for following the majority class."
    )
    add_figure(doc, "task4_crime_type_distribution.png", "Figure 6 Distribution of the most frequent Primary Type categories")
    add_paragraph(doc,
        "The feature distribution is also skewed. Theft represents 21.16 per cent of modelling records, battery 18.33 per cent, criminal damage 11.45 per cent and narcotics 9.43 per cent. Less frequent crime categories exert less influence during tree construction. More importantly, arrest practices differ strongly by offence type, so Primary Type can dominate prediction. Feature importance confirms this: Primary Type contributes 0.937 of total impurity-based importance, while hour contributes 0.023, latitude 0.018, longitude 0.012, community area 0.007 and district 0.002. Month has zero importance in the fitted model. This concentration improves predictive fit but means the model largely learns historical links between recorded offence category and arrest outcome."
    )

    add_heading(doc, "6.2 Validation design and evaluation measures", 2)
    add_paragraph(doc,
        "The main comparison uses a seeded 80 to 20 random holdout. Inverse-frequency weights give each class equal total training weight. Unweighted and weighted Random Forests establish the effect of imbalance (Breiman, 2001). Weighted Logistic Regression provides a linear comparison with one-hot-encoded Primary Type, while weighted Gradient Boosted Trees provide a sequential nonlinear ensemble (Friedman, 2001). All models use the same split and input concepts. BinaryClassificationEvaluator calculates ROC AUC and PR AUC (Apache Software Foundation, 2026a), while Spark aggregations produce the confusion counts. F1 combines precision and recall, and balanced accuracy averages recall with specificity."
    )
    add_paragraph(doc,
        "PR AUC is included because ROC performance can look favourable when the negative class is much larger. Saito and Rehmsmeier (2015) show that precision-recall analysis is especially informative for imbalanced binary data. The evaluation therefore avoids relying on one metric. High precision can coexist with weak recall, while a lower threshold can increase recall at the cost of more false positives. The appropriate trade-off depends on use. In this project, predictions are analytical evidence rather than a basis for enforcement, so no operational cost ratio is claimed."
    )
    add_table(doc,
        ["Model and validation", "Precision", "Recall", "F1", "ROC AUC", "PR AUC"],
        [
            ["Majority baseline random holdout", "0.000", "0.000", "0.000", "0.500", "0.258"],
            ["Unweighted RF random holdout", "0.994", "0.362", "0.531", "0.813", "0.745"],
            ["Weighted RF random holdout", "0.941", "0.461", "0.619", "0.818", "0.740"],
            ["Weighted logistic random holdout", "0.886", "0.558", "0.685", "0.834", "0.765"],
            ["Weighted GBT random holdout", "0.831", "0.582", "0.684", "0.848", "0.778"],
            ["Weighted RF temporal holdout", "0.759", "0.365", "0.493", "0.795", "0.553"],
        ],
        [2.55, 0.85, 0.8, 0.75, 0.9, 0.85],
        "Table 6 Arrest model comparison",
    )
    add_paragraph(doc,
        "The comparison shows that Random Forest was not the strongest predictive model. Weighted Logistic Regression produced the highest F1 score at 0.685, with precision of 0.886 and recall of 0.558. Weighted Gradient Boosted Trees achieved almost the same F1 score at 0.684 but the strongest recall of 0.582, balanced accuracy of 0.770, ROC AUC of 0.848 and PR AUC of 0.778. Random Forest remains useful because it provides compact feature-importance evidence and supports clear threshold analysis. Logistic Regression offers the simpler decision function, whereas boosting gains ranking performance through sequential nonlinear trees (Friedman, 2001). These results show why algorithm selection should consider several metrics rather than familiarity with one model."
    )

    add_heading(doc, "6.3 Model results and threshold analysis", 2)
    add_figure(doc, "task4_model_comparison.png", "Figure 7 Precision recall and F1 across random holdout models")
    add_paragraph(doc,
        "The unweighted Random Forest achieved precision of 0.994 but recall of only 0.362. It made 132,752 correct arrest predictions and only 839 false-positive predictions, yet it missed 233,852 arrests. The result illustrates the effect of class skew: the model is extremely cautious when predicting the minority class. Its F1 score was 0.531 and balanced accuracy was 0.681. ROC AUC of 0.813 and PR AUC of 0.745 show useful ranking performance, but the default class decision did not convert that ranking into adequate recall."
    )
    add_paragraph(doc,
        "Class weighting changed the balance. The weighted model produced 168,949 true positives, 10,581 false positives, 197,655 false negatives and 1,042,428 true negatives. Recall rose to 0.461 and F1 to 0.619, while precision declined to 0.941. Balanced accuracy improved to 0.725 and ROC AUC to 0.818. PR AUC was slightly lower at 0.740. Within the two Random Forest variants, the weighted model is preferred because it identifies 36,197 additional arrests and achieves a much stronger F1 score, although it also creates more false positives. This is a measured trade-off rather than a claim that weighting improves every metric."
    )
    add_figure(doc, "task4_roc_curve.png", "Figure 8 Receiver operating characteristic curve for the weighted Random Forest")
    add_figure(doc, "task4_threshold_metrics.png", "Figure 9 Precision recall and F1 across probability thresholds")
    add_paragraph(doc,
        "Threshold analysis makes the trade-off explicit. At 0.2, recall is 1.0 but precision falls to 0.258 because nearly every record is classified as an arrest. At 0.4, recall is 0.763, precision is 0.455 and F1 is 0.570. The default 0.5 threshold gives the highest tested F1 score of 0.619, with precision of 0.941 and recall of 0.461. At 0.6, precision rises to 0.994 while recall returns to 0.362. A threshold should therefore be selected from a stated objective and cost of error. Maximising recall without considering false positives would be inappropriate in a policing context, while maximising precision alone would miss most positive cases."
    )
    add_figure(doc, "task4_feature_importance.png", "Figure 10 Random Forest feature importance")
    add_paragraph(doc,
        "A second validation trains on records before 2020 and tests on 2020 onwards. The temporal holdout is more demanding because it measures performance on later conditions rather than a random sample from the same overall period. Precision fell to 0.759, recall to 0.365, F1 to 0.493, ROC AUC to 0.795 and PR AUC to 0.553. This gap indicates temporal drift or changes in the relationship between incident characteristics and arrests. Possible causes include changes in crime composition, recording, policy and police practice. The temporal result is more relevant to future deployment and shows that the random holdout alone was optimistic."
    )
    add_heading(doc, "7 Performance reproducibility and software quality", 1)
    add_heading(doc, "7.1 Performance decisions", 2)
    add_paragraph(doc,
        "Performance improvements were applied where they matched the data shape. Parquet avoids repeated CSV parsing and supports column pruning. Year and district partitions support selective reads. The census table is broadcast because it is tiny relative to the crime data. Crime counts are aggregated to hour before joining weather, reducing join volume. K selection uses a reproducible sample, while the chosen model still assigns all valid crime records. Modelling coalesces the source to eight partitions and persists reused frames on disk. These choices reduce unnecessary shuffles and local scheduling overhead without changing the analytical definitions."
    )
    add_paragraph(doc,
        "Some optimisation opportunities remain. The Parquet crime output contains many year and district combinations, and repeated overwrite operations can create small files. A production workflow should monitor file counts and target file size, then compact partitions when required. The global ROC curve orders score groups in one window and Spark warns that a window without a partition moves data to a single partition. The Random Forest produces a limited number of score values, so this completed locally, but a model with many unique scores would need an approximate or binned ROC calculation. These observations show that successful execution does not remove the need to inspect physical plans and runtime warnings."
    )
    add_paragraph(doc,
        "The local run used four workers, eight shuffle partitions and a two-gigabyte driver. Spark spilled several Random Forest blocks to disk because they did not fit in memory. Disk spill is slower but preferable to driver failure. A multi-node deployment should distribute training partitions across executors, separate driver and executor memory settings, and store Parquet on distributed or object storage. It should also collect run time, shuffle size, spill volume and task skew before tuning. The present project documents the local constraint rather than claiming performance numbers that would generalise to a cluster."
    )

    add_heading(doc, "7.2 Reproducibility and testing", 2)
    add_paragraph(doc,
        "Reproducibility is supported through fixed random seeds, exact dependency versions, command-line paths and stored analytical outputs. The requirements file pins PySpark 3.5.9, NumPy 2.5.3, pandas 2.3.3, Matplotlib 3.11.2, pytest 8.4.2 and setuptools 80.10.2. Each task writes Parquet tables and report-ready CSV evidence. Trained model files and charts are also retained. A reader can therefore inspect the numerical outputs without rerunning the 1.7-gigabyte source file, while the commands remain available when a full rerun is required."
    )
    add_paragraph(doc,
        "The test suite contains thirteen passing tests. Transformation tests cover invalid timestamps, duplicate IDs, null IDs and geographic flags. Weather tests cover trace precipitation, event parsing, hourly aggregation and the left join that fills missing crime hours with zero. Modelling tests verify ROC endpoints, equal total class weight, threshold trade-offs and the comparison pipeline estimators. An integration test writes cleaned records partitioned by year and district, reads them back and confirms the directory structure. These tests protect the key business rules and one end-to-end storage path."
    )
    add_paragraph(doc,
        "Clean code practices improve the credibility of the analysis because they reduce accidental inconsistencies. Dataclasses group related outputs, functions have descriptive names and type hints, and technical docstrings explain public interfaces. Shared expressions and pipeline builders follow the do not repeat yourself principle. Logging replaces scattered terminal print statements, except where Spark DataFrame display is useful for compact run summaries. The design remains simple and does not introduce abstract classes or service layers where functions and immutable configuration are sufficient. SOLID principles are applied in proportion to the project rather than used to justify unnecessary complexity."
    )

    add_heading(doc, "8 Ethical issues and study limitations", 1)
    add_heading(doc, "8.1 Meaning of reported crime and arrest labels", 2)
    add_paragraph(doc,
        "The dataset measures contact with reporting and policing systems, not all harmful events. An incident can be absent because it was not reported, not recorded or later classified differently. The City of Chicago warns that information may be preliminary and can contain human or mechanical error (City of Chicago, 2026). Arrest is also an institutional outcome. It depends on evidence, police presence, enforcement priorities and legal processes, as well as incident characteristics. A model that predicts arrest therefore learns historical practice rather than a neutral property of an offence."
    )
    add_paragraph(doc,
        "This distinction is especially important because location and crime category are prominent features. Lum and Isaac (2016) argue that police databases reflect an interaction among criminal activity, policing strategy and community police relationships, and show how predictive policing can reproduce patterns in its training data. The current model should not be used to allocate patrols, label communities or influence decisions about individuals. Its purpose is to demonstrate distributed classification and evaluate error patterns. Any proposed operational use would require legal review, community involvement, fairness analysis and model governance. A model card should also document intended use, evaluation conditions and known limitations before release (Mitchell et al., 2019)."
    )

    add_heading(doc, "8.2 Analytical limitations", 2)
    add_paragraph(doc,
        "Several limitations affect Task 2. Pearson correlation measures linear association and does not control for confounding. Seasonal tables improve interpretation but do not replace a multivariable model. Daily crime counts are not adjusted for population or the number of people present in each area. Community hardship values cover 2008 to 2012 and should not be applied to later years as if they were constant. Airport weather does not capture neighbourhood variation. Recorded incident time may differ from occurrence time, particularly when midnight is used as a default. These limits mean the analysis identifies patterns worth investigating, not causal relationships."
    )
    add_paragraph(doc,
        "K-means seeks compact groups under squared Euclidean distance and requires K in advance. Geographic coordinates are treated on a flat numerical plane rather than a true distance projection, and hour is linear rather than circular. Three seeds measure sensitivity to initialisation, but one training sample does not measure sampling or temporal stability. The two selected clusters cross 22 districts each and are too broad for direct resource planning. Density-based clustering, projected coordinates, cyclic time and repeated samples would provide a stronger hotspot study."
    )
    add_paragraph(doc,
        "The predictive models contain no narrative, victim, officer, case progression or contextual variables. This reduces privacy concerns but limits recall. Random splitting allows records from similar periods on both sides and therefore overstates future performance compared with the temporal holdout. Three algorithm families are compared, but each uses one main parameter setting and only the Random Forest receives temporal evaluation. Probabilities are not calibrated, and error rates are not assessed across demographic groups because protected characteristics are absent and should not be inferred from neighbourhood. The models should remain coursework evidence rather than decision support."
    )

    add_heading(doc, "9 Conclusion and recommendations", 1)
    add_paragraph(doc,
        "The project completed all four analytical tasks in one reproducible Spark codebase. Task 1 converted 7,784,664 raw crime records into typed Parquet without losing valid incidents, reported geographic defects explicitly, created a quarantine path and produced a data dictionary. Census and weather sources were cleaned at appropriate levels of detail. Task 2 used Spark SQL, windows and broadcast joins to identify temporal, community and weather patterns. The temperature correlation of 0.510 and the 12.8 per cent lower count on snow days were descriptive findings, while seasonal comparison showed why neither should be interpreted as a simple causal effect."
    )
    add_paragraph(doc,
        "Task 3 selected K equals two through a mean silhouette of 0.496. Its three-seed standard deviation of 0.00035 showed stable initialisation, but both broad clusters crossed all 22 districts. Task 4 showed that algorithm choice mattered. Logistic Regression achieved the highest F1 score of 0.685, while Gradient Boosted Trees achieved the best recall of 0.582, balanced accuracy of 0.770, ROC AUC of 0.848 and PR AUC of 0.778. The Random Forest temporal F1 score of 0.493 provided direct evidence of drift and showed why random holdout performance should not be treated as future performance."
    )
    add_paragraph(doc,
        "The strongest next improvements are methodological. Weather analysis should use several stations and a multivariable model with season, weekday, year and exposure controls. Spatial analysis should use projected coordinates, cyclic time and repeated samples, then compare K-means with a density-based method. Arrest modelling should add calibration, temporal cross-validation for every candidate and systematic hyperparameter tuning. Fairness work must use appropriately governed data and examine institutional processes rather than infer sensitive characteristics from location. Performance testing should move to a multi-node environment and record shuffle, spill and file statistics."
    )
    add_paragraph(doc,
        "Overall, Apache Spark was an appropriate platform for this workload because it combined typed distributed transformations, SQL, window analysis, broadcast joins, Parquet storage and machine learning within one execution model. The project also shows that technical scale does not guarantee analytical validity. High-quality work requires traceable quality decisions, comparison baselines, validation that reflects future use, and clear limits on what the data can support. These principles make the results more credible and provide a sound basis for extending the system."
    )

    add_heading(doc, "References", 1)
    references = [
        "Apache Parquet (2026) Apache Parquet. Available at: https://parquet.apache.org/ (Accessed: 23 September 2026).",
        "Apache Software Foundation (2026a) BinaryClassificationEvaluator: Spark 3.5.9 ScalaDoc. Available at: https://spark.apache.org/docs/3.5.9/api/scala/org/apache/spark/ml/evaluation/BinaryClassificationEvaluator.html (Accessed: 23 September 2026).",
        "Apache Software Foundation (2026b) Clustering: Spark 3.5.9 documentation. Available at: https://spark.apache.org/docs/3.5.9/ml-clustering.html (Accessed: 23 September 2026).",
        "Apache Software Foundation (2026c) Data sources: Spark 3.5.9 documentation. Available at: https://spark.apache.org/docs/3.5.9/sql-data-sources.html (Accessed: 23 September 2026).",
        "Apache Software Foundation (2026d) Performance tuning: Spark 3.5.9 documentation. Available at: https://spark.apache.org/docs/3.5.9/sql-performance-tuning.html (Accessed: 23 September 2026).",
        "Apache Software Foundation (2026e) Window functions: Spark 3.5.9 documentation. Available at: https://spark.apache.org/docs/3.5.9/sql-ref-syntax-qry-select-window.html (Accessed: 23 September 2026).",
        "Bahmani, B., Moseley, B., Vattani, A., Kumar, R. and Vassilvitskii, S. (2012) Scalable k-means++. Proceedings of the VLDB Endowment, 5(7), pp. 622-633. doi: 10.14778/2180912.2180915.",
        "Breiman, L. (2001) Random forests. Machine Learning, 45, pp. 5-32. doi: 10.1023/A:1010933404324.",
        "Chicago Department of Public Health (2014) Selected socioeconomic indicators in Chicago 2008 to 2012 Dataset Description. Available at: https://data.cityofchicago.org/api/views/kn9c-c2s2/ (Accessed: 23 September 2026).",
        "City of Chicago (2026) Crimes 2001 to Present. Available at: https://data.cityofchicago.org/Public-Safety/Crimes-2001-to-Present/ijzp-q8t2 (Accessed: 23 September 2026).",
        "Friedman, J.H. (2001) Greedy function approximation: A gradient boosting machine. The Annals of Statistics, 29(5), pp. 1189-1232. doi: 10.1214/aos/1013203451.",
        "He, H. and Garcia, E.A. (2009) Learning from imbalanced data. IEEE Transactions on Knowledge and Data Engineering, 21(9), pp. 1263-1284. doi: 10.1109/TKDE.2008.239.",
        "Kantor, D., Casey, N.W., Menne, M.J. and Buddenberg, A. (2023) Local Climatological Data Version 2. NOAA National Centers for Environmental Information. Available at: https://www.ncei.noaa.gov/products/land-based-station/local-climatological-data (Accessed: 23 September 2026).",
        "Lum, K. and Isaac, W. (2016) To predict and serve. Significance, 13(5), pp. 14-19. doi: 10.1111/j.1740-9713.2016.00960.x.",
        "Mitchell, M., Wu, S., Zaldivar, A., Barnes, P., Vasserman, L., Hutchinson, B., Spitzer, E., Raji, I.D. and Gebru, T. (2019) Model cards for model reporting. Proceedings of the Conference on Fairness Accountability and Transparency, pp. 220-229. doi: 10.1145/3287560.3287596.",
        "Rousseeuw, P.J. (1987) Silhouettes: A graphical aid to the interpretation and validation of cluster analysis. Journal of Computational and Applied Mathematics, 20, pp. 53-65. doi: 10.1016/0377-0427(87)90125-7.",
        "Saito, T. and Rehmsmeier, M. (2015) The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. PLOS ONE, 10(3), e0118432. doi: 10.1371/journal.pone.0118432.",
        "Zaharia, M., Xin, R.S., Wendell, P., Das, T., Armbrust, M., Dave, A., Meng, X., Rosen, J., Venkataraman, S., Franklin, M.J., Ghodsi, A., Gonzalez, J., Shenker, S. and Stoica, I. (2016) Apache Spark: A unified engine for big data processing. Communications of the ACM, 59(11), pp. 56-65. doi: 10.1145/2934664.",
    ]
    for reference in references:
        p = add_paragraph(doc, reference, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
        p.paragraph_format.left_indent = Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(-1.27)

    add_heading(doc, "Appendix A Reproducibility commands", 1)
    commands = [
        "PYTHONPATH=src .venv/bin/python -m chicago_crime.run_pipeline --input data/raw/chicago_crime.csv --output data/processed/chicago_crime_parquet --quality-output data/processed/task1_crime_quality --quarantine-output data/processed/task1_quarantine --dictionary-output data/processed/task1_data_dictionary",
        "PYTHONPATH=src .venv/bin/python -m chicago_crime.task1 --census-input data/raw/chicago_census.csv --census-output data/processed/chicago_census_parquet --weather-input data/raw/chicago_weather.csv --weather-output data/processed/chicago_weather_parquet --weather-quality-output data/processed/task1_weather_quality",
        "PYTHONPATH=src .venv/bin/python -m chicago_crime.run_tasks eda --crime-input data/processed/chicago_crime_parquet --weather-input data/processed/chicago_weather_parquet --census-input data/processed/chicago_census_parquet --output-dir outputs",
        "PYTHONPATH=src .venv/bin/python -m chicago_crime.run_tasks cluster --crime-input data/processed/chicago_crime_parquet --output-dir outputs",
        "PYTHONPATH=src .venv/bin/python -m chicago_crime.run_tasks model --crime-input data/processed/chicago_crime_parquet --output-dir outputs",
        "PYTHONPATH=src .venv/bin/python -m pytest -q",
    ]
    for command in commands:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.right_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(6)
        set_run_font(p.add_run(command), name="Courier New", size=8)

    core_text = []
    in_core = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        is_heading = paragraph.style.name.startswith("Heading")
        if is_heading and text in {"Abstract", "1 Introduction"}:
            in_core = True
        if is_heading and text == "Contents":
            in_core = False
        if is_heading and text == "References":
            in_core = False
        if in_core and text not in {"Abstract", "Contents"} and not re.match(r"^\d+(\.\d+)?\s", text):
            core_text.append(text)
    word_count = len(re.findall(r"\b[\w]+(?:['’-][\w]+)*\b", " ".join(core_text)))
    for table in doc.tables[:1]:
        for row in table.rows:
            if row.cells[0].text == "Word count":
                row.cells[1].text = f"{word_count:,} words excluding references and appendices"
                for run in row.cells[1].paragraphs[0].runs:
                    set_run_font(run)
    properties = doc.core_properties
    properties.title = "Distributed Analysis of Chicago Crime Data Using Apache Spark"
    properties.subject = "DAS 7002 Big Data Technologies PRAC 1"
    properties.keywords = "Apache Spark, PySpark, Chicago crime, K-means, classification"
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Core word count: {word_count}")


if __name__ == "__main__":
    build_document()
